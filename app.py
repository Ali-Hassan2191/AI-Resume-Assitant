import json
import os
import re
import tempfile
from pathlib import Path

import streamlit as st
from google import genai
from google.genai import types
from pypdf import PdfReader
from docx import Document


st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

MODEL_NAME = "gemini-3.6-flash"

def get_api_key():
    """Read Gemini API key from Streamlit secrets or environment."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
        if key:
            return key
    except Exception:
        pass

    return os.getenv("GEMINI_API_KEY")


def extract_docx_text(uploaded_file):
    """Extract readable text from a DOCX resume."""
    data = uploaded_file.getvalue()
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(data)
        tmp_path = tmp.name

    try:
        document = Document(tmp_path)
        parts = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    parts.append(row_text)

        return "\n".join(parts).strip()
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def extract_pdf_text(uploaded_file):
    """Extract text from a PDF for basic ATS checks."""
    reader = PdfReader(uploaded_file)
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)

    return "\n".join(pages).strip()


def basic_ats_checks(resume_text):
    """
    Deterministic checks provide a baseline. Gemini performs the deeper analysis.
    """
    text = resume_text.strip()
    lower = text.lower()

    sections = {
        "contact": any(
            item in lower
            for item in ["@", "phone", "linkedin", "github", "contact"]
        ),
        "summary": any(
            item in lower
            for item in ["summary", "objective", "profile", "professional summary"]
        ),
        "experience": any(
            item in lower
            for item in ["experience", "work experience", "employment", "professional experience"]
        ),
        "education": "education" in lower or "academic" in lower,
        "skills": "skills" in lower or "technical skills" in lower,
    }

    bullet_count = len(re.findall(r"(?m)^\s*(?:[-•▪◦*]|\d+[.)])\s+", text))
    word_count = len(re.findall(r"\b[\w+#.-]+\b", text))

    return {
        "word_count": word_count,
        "bullet_count": bullet_count,
        "sections_found": sections,
    }


def analyze_resume(uploaded_file, job_description):
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError(
            "Gemini API key is missing. Add GEMINI_API_KEY to Streamlit Secrets "
            "or your local environment."
        )

    client = genai.Client(api_key=api_key)
    file_type = uploaded_file.type or ""
    file_name = uploaded_file.name.lower()

    baseline = {}

    if file_type == "application/pdf" or file_name.endswith(".pdf"):
        # Gemini can inspect PDFs natively, which is useful for resume layout/formatting.
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(uploaded_file.getvalue())
            tmp_path = tmp.name

        try:
            gemini_file = client.files.upload(
                file=tmp_path,
                config={"mime_type": "application/pdf"},
            )

            prompt = build_prompt(
                job_description=job_description,
                extracted_text=None,
                baseline=None,
            )

            response = client.models.generate_content(
                model=MODEL_NAME,
                contents=[gemini_file, prompt],
                config=types.GenerateContentConfig(
                    temperature=0.2,
                    max_output_tokens=5000,
                    response_mime_type="application/json",
                ),
            )
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    elif (
        file_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or file_name.endswith(".docx")
    ):
        extracted_text = extract_docx_text(uploaded_file)
        if not extracted_text:
            raise ValueError("The DOCX file does not contain readable text.")

        baseline = basic_ats_checks(extracted_text)
        prompt = build_prompt(
            job_description=job_description,
            extracted_text=extracted_text,
            baseline=baseline,
        )

        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=prompt,
            config=types.GenerateContentConfig(
                temperature=0.2,
                max_output_tokens=5000,
                response_mime_type="application/json",
            ),
        )

    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX resume.")

    raw = response.text.strip()
    return parse_json_response(raw)


def build_prompt(job_description, extracted_text=None, baseline=None):
    job_text = job_description.strip() if job_description else ""

    source = ""
    if extracted_text:
        source = f"""
RESUME TEXT:
{extracted_text}
"""
        source += f"""
DETERMINISTIC BASELINE CHECKS:
{json.dumps(baseline, indent=2)}
"""
    else:
        source = """
The uploaded document is the resume itself. Inspect both its content and visible
PDF layout. Do not invent information that is not present in the resume.
"""

    if job_text:
        job_section = f"""
TARGET JOB DESCRIPTION:
{job_text}

Use the job description for keyword and relevance matching. Separate exact/missing
keywords from generic resume advice.
"""
    else:
        job_section = """
No target job description was supplied. Evaluate ATS compatibility using general
ATS-friendly resume practices, while clearly stating that keyword matching is
general rather than job-specific.
"""

    return f"""
You are an expert resume reviewer and ATS compatibility analyst.

Analyze the supplied resume and return ONLY valid JSON.

Important:
- The ATS score is an ESTIMATED compatibility score, not a score from a specific
  commercial ATS vendor.
- Do not judge the person, age, gender, race, religion, nationality, disability,
  or other protected/personal characteristics.
- Do not invent work experience, education, skills, dates, employers, or metrics.
- Give practical improvements that the candidate can actually make.
- Consider ATS parsing, section structure, standard headings, keyword alignment,
  readability, measurable achievements, consistency, and formatting risks.
- If the PDF contains tables, columns, icons, graphics, headers/footers, or other
  layout elements that may create parsing risk, mention them only when you can
  actually observe them.
{source}
{job_section}

Return exactly this JSON structure:
{{
  "ats_score": 0,
  "score_label": "Needs improvement",
  "score_explanation": "Short explanation of why the estimated score was given.",
  "score_breakdown": {{
    "parseability": 0,
    "keyword_alignment": 0,
    "section_structure": 0,
    "skills_relevance": 0,
    "achievement_quality": 0
  }},
  "summary": "2-4 sentence overall resume assessment.",
  "strengths": [
    "strength 1",
    "strength 2",
    "strength 3"
  ],
  "ats_risks": [
    "specific ATS risk 1",
    "specific ATS risk 2",
    "specific ATS risk 3"
  ],
  "critical_improvements": [
    {{
      "issue": "Specific issue",
      "why_it_matters": "Why this can hurt ATS parsing or recruiter readability.",
      "fix": "Specific action to take."
    }}
  ],
  "section_feedback": [
    {{
      "section": "Experience",
      "status": "Good",
      "feedback": "Specific feedback."
    }}
  ],
  "keywords": {{
    "matched": ["keyword 1"],
    "missing_or_weak": ["keyword 2"],
    "notes": "Explain keyword matching briefly."
  }},
  "formatting_risks": [
    "risk 1"
  ],
  "recommended_resume_structure": [
    "Contact Information",
    "Professional Summary",
    "Skills",
    "Professional Experience",
    "Education"
  ],
  "rewrites": [
    {{
      "original": "A short existing bullet from the resume, or a concise description of the weak bullet.",
      "improved": "An improved version that does not invent facts."
    }}
  ],
  "next_steps": [
    "Action 1",
    "Action 2",
    "Action 3"
  ]
}}

Scoring guidance:
- 90-100: Very strong ATS compatibility
- 75-89: Good, with some improvements
- 60-74: Moderate; several improvements recommended
- 0-59: Significant ATS/readability improvements needed

For score_breakdown, return an independent 0-100 score for each:
- parseability: How reliably an ATS can extract and search the resume content.
- keyword_alignment: How well the resume matches the target job description, or general ATS keywords when no job description is provided.
- section_structure: Quality and clarity of standard resume sections and headings.
- skills_relevance: Relevance and strength of the listed skills for the target role.
- achievement_quality: Strength of experience/project bullets, action verbs, measurable results, and evidence of impact.

Base every sub-score only on evidence in the resume and supplied job description.
Make ats_score consistent with the five breakdown scores by using their rounded average.
Return concise ATS risks that are specific to the uploaded resume.

Return no Markdown fences and no text outside the JSON.
"""


def parse_json_response(raw):
    """Handle occasional markdown fences even when JSON mode is requested."""
    cleaned = raw.strip()

    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "Gemini returned an invalid analysis format. Please try again."
        ) from exc

    score = data.get("ats_score")
    if not isinstance(score, (int, float)):
        raise ValueError("Gemini did not return a valid ATS score.")

    data["ats_score"] = max(0, min(100, int(round(score))))
    return data


def _score_status(score):
    if score >= 80:
        return "Strong"
    if score >= 60:
        return "Fair"
    return "Needs improvement"


def display_results(result):
    score = result["ats_score"]
    label = result.get("score_label", "")

    # Overall score header
    st.subheader("ATS Compatibility Score")
    st.progress(score / 100)
    col1, col2 = st.columns([1, 3])
    with col1:
        st.metric("Estimated ATS Score", f"{score}/100")
    with col2:
        st.write(f"**{label}**")
        st.write(result.get("score_explanation", ""))

    # Score Breakdown — added to match the requested reference UI.
    st.subheader("📊 Score Breakdown")
    breakdown = result.get("score_breakdown", {})
    score_items = [
        ("Parseability", "parseability"),
        ("Keyword Alignment", "keyword_alignment"),
        ("Section Structure", "section_structure"),
        ("Skills Relevance", "skills_relevance"),
        ("Achievement Quality", "achievement_quality"),
    ]

    cols = st.columns(5)
    for col, (title, key) in zip(cols, score_items):
        try:
            value = int(round(float(breakdown.get(key, 0))))
        except (TypeError, ValueError):
            value = 0
        value = max(0, min(100, value))
        with col:
            st.metric(title, f"{value}/100")
            st.caption(_score_status(value))

    st.subheader("Overall Assessment")
    st.write(result.get("summary", ""))

    # Strengths + ATS Risks side by side — matching the reference image.
    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        strengths = result.get("strengths", [])
        if strengths:
            for item in strengths:
                st.markdown(f"- {item}")
        else:
            st.write("No strengths were returned.")

    with right:
        st.subheader("⚠️ ATS Risks")
        risks = result.get("ats_risks")
        if risks is None:
            risks = result.get("formatting_risks", [])
        if risks:
            for risk in risks:
                st.markdown(f"- {risk}")
        else:
            st.success("No major ATS risks were identified.")

    st.subheader("🎯 Critical Improvements")
    for item in result.get("critical_improvements", []):
        if isinstance(item, dict):
            st.markdown(f"**{item.get('issue', 'Issue')}**")
            st.write(item.get("why_it_matters", ""))
            st.info(f"Fix: {item.get('fix', '')}")
        else:
            st.markdown(f"- {item}")

    st.subheader("Section-by-Section Feedback")
    for item in result.get("section_feedback", []):
        section = item.get("section", "Section")
        status = item.get("status", "")
        feedback = item.get("feedback", "")
        with st.expander(f"{section} — {status}"):
            st.write(feedback)

    st.subheader("🔎 Keyword Analysis")
    keywords = result.get("keywords", {})
    k1, k2 = st.columns(2)
    with k1:
        st.markdown("**Matched / present**")
        matched = keywords.get("matched", [])
        st.write(", ".join(matched) if matched else "No clear matches identified.")
    with k2:
        st.markdown("**Missing / weak**")
        missing = keywords.get("missing_or_weak", [])
        st.write(", ".join(missing) if missing else "No major missing keywords identified.")
    if keywords.get("notes"):
        st.caption(keywords["notes"])

    st.subheader("✍️ Suggested Bullet Improvements")
    rewrites = result.get("rewrites", [])
    if rewrites:
        for item in rewrites:
            with st.expander("View improvement"):
                st.markdown("**Current:**")
                st.write(item.get("original", ""))
                st.markdown("**Improved:**")
                st.write(item.get("improved", ""))
    else:
        st.write("No specific bullet rewrites were returned.")

    st.subheader("Recommended Resume Structure")
    for item in result.get("recommended_resume_structure", []):
        st.markdown(f"- {item}")

    st.subheader("🚀 Next Steps")
    for index, item in enumerate(result.get("next_steps", []), start=1):
        st.markdown(f"{index}. {item}")

st.title("📄 Resume ATS Analyzer")
st.caption(
    "Upload a resume to get an estimated ATS compatibility score, keyword analysis, "
    "formatting risks, and practical improvement suggestions."
)

with st.sidebar:
    st.header("About")
    st.write(
        "This tool uses Gemini Flash to analyze resume content and, for PDFs, "
        "the document's visual layout."
    )
    st.caption(
        "The score is an AI-based estimate. Different ATS platforms can parse "
        "the same resume differently."
    )

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="Supported formats: PDF and DOCX. PDF is recommended for layout analysis.",
)

job_description = st.text_area(
    "Optional: paste the target job description",
    height=220,
    placeholder=(
        "Paste the job posting here to get job-specific keyword matching and "
        "relevance feedback."
    ),
)

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:
    if not uploaded_file:
        st.warning("Please upload a PDF or DOCX resume first.")
    else:
        with st.spinner("Analyzing your resume with Gemini Flash..."):
            try:
                result = analyze_resume(uploaded_file, job_description)
                st.session_state["analysis_result"] = result
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")

if "analysis_result" in st.session_state:
    display_results(st.session_state["analysis_result"])
